# 自動填寫大宗郵件文件
# 準備好src.xlsx (A column填姓名 B column填三碼郵遞區號或地址)
# (範例姓名地址皆為虛構)
# 依照post.docx格式輸出
# By StevenYU 2018/01/25
# run at python3

# coding=UTF-8
from docx import Document
from openpyxl import load_workbook
import math

wb = load_workbook(filename=r'src.xlsx')
sh = wb['src_data']
s = int(input("資料數: ")) + 1
name = []
zip_code = []
for i in range(1,s):
    name.append(str(sh[('A'+str(i))].value))
    zip_code.append((str(sh[('B'+str(i))].value))[:3])
doc = Document("post.docx")

counter = 0

for total in range(1,int(math.ceil(s/20)+1)):
    for table in doc.tables:
        i = 0
        counter = (total-1)*20
        for col in table.column_cells(2):
            if i<=1 :
                i+=1
                continue
            try:
                col.text = name[counter]
            except:
                col.text = ""
            counter += 1

        i = 0
        counter = (total-1)*20
        for col in table.column_cells(3):
            if i<=1 :
                i+=1
                continue
            try:
                col.text = zip_code[counter]
            except:
                col.text = ""
            counter += 1

        doc.save("out"+ str(total)+".docx")
